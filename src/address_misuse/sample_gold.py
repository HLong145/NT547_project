"""D5 — lấy mẫu gán nhãn vàng (pilot 50 mẫu) kèm bằng chứng điền sẵn, và tính Cohen's κ.

python -X utf8 -m address_misuse.sample_gold pilot [--seed 0]
python -X utf8 -m address_misuse.sample_gold main [--seed 0] [--double 200]   → mẫu vàng chính từ dataset_v2 (sau freeze)
python -X utf8 -m address_misuse.sample_gold sheet B      → data/gold/annotator_B.xlsx (file Excel để gán)
python -X utf8 -m address_misuse.sample_gold sheet B --set main [--only-double | --only-single]   → annotator_main[_double|_single]_B.xlsx
python -X utf8 -m address_misuse.sample_gold check data/gold/annotator_B.xlsx   (kiểm tra hình thức trước khi tính κ)
python -X utf8 -m address_misuse.sample_gold kappa data/gold/annotator_A.csv data/gold/annotator_B.csv

Đầu ra trong data/gold/ (chỉ dùng nội bộ, có địa chỉ đầy đủ):
- pilot50_items.csv       danh sách đã làm mù (không có tầng/nhãn tác giả), thứ tự xáo trộn, kèm bằng chứng
- pilot50_key.csv         tầng lấy mẫu + nhãn tác giả (RQ1) — người gán KHÔNG mở file này
- annotator_template.csv  cột cần điền theo docs/CODEBOOK.md §2
Định nghĩa nhãn (cho người gán): docs/CODEBOOK.md. Thiết kế mẫu, quy trình, báo cáo, hạn chế: docs/GOLD_PROTOCOL.md.
"""

import argparse
import json
from datetime import datetime, timezone

import numpy as np
import pandas as pd

from address_misuse.config import GOLD_DIR, INTERIM_DIR, NETWORKS, RAW_DIR, SEED_DIR

EXPLORERS = {
    "eth": "https://etherscan.io/address/", "bsc": "https://bscscan.com/address/",
    "arb": "https://arbiscan.io/address/", "op": "https://optimistic.etherscan.io/address/",
    "base": "https://basescan.org/address/", "polygon": "https://polygonscan.com/address/",
    "sepolia": "https://sepolia.etherscan.io/address/", "bsc_testnet": "https://testnet.bscscan.com/address/",
}
LABEL_COLS = ["is_misuse", "type", "subtype", "xc_chain", "malicious", "confidence", "evidence", "note"]
SYMEXEC_FLAGS = ["delegatecall", "transfer_all_eth", "transfer_msg_value", "transfer_input_amount"]

# (tầng, số mẫu). Tổng 50. Ưu tiên Type 1/4/5 (lõi) và có đủ âm để kiểm tra codebook.
PILOT_PLAN = [
    ("rq1_ca_pos", 3), ("rq1_ca_neg", 2), ("rq1_eoa_addr", 3),     # kiểm tra lại nhãn tác giả
    ("t1_ncd_other", 10),                                          # NCD không có contract ở chain khác
    ("t4_xc_before", 3), ("t4_xc_bsc", 3), ("t4_xc_after", 2), ("t4_xc_unknown", 2),
    ("ca_pcd", 4),
    ("eoa_key", 4), ("eoa_addr", 2), ("eoa_7702", 2), ("eoa_7702_other", 2),
    ("neg_n7_multichain", 4), ("neg_n1_random", 4),
]


def _date(ts) -> str | None:
    if ts is None or (isinstance(ts, float) and np.isnan(ts)) or pd.isna(ts):
        return None
    return datetime.fromtimestamp(int(float(ts)), tz=timezone.utc).strftime("%Y-%m-%d")


def _dedup_max(df: pd.DataFrame, keys=("chain", "address")) -> pd.DataFrame:
    return df.sort_values("cost_native", ascending=False).drop_duplicates(list(keys))


def _pick(df: pd.DataFrame, n: int, rng: np.random.Generator, loss_col: str | None = "cost_native") -> pd.DataFrame:
    """Nửa từ nhóm thiệt hại cao (top 25%), nửa ngẫu nhiên phần còn lại — để có cả ca lớn lẫn ca điển hình."""
    if len(df) <= n:
        return df
    if not loss_col or loss_col not in df:
        return df.sample(n, random_state=rng.integers(2**31))
    q = df[loss_col].quantile(0.75)
    hi, lo = df[df[loss_col] >= q], df[df[loss_col] < q]
    n_hi = min(len(hi), n // 2)
    a = hi.sample(n_hi, random_state=rng.integers(2**31))
    b = lo.sample(min(len(lo), n - n_hi), random_state=rng.integers(2**31))
    out = pd.concat([a, b])
    if len(out) < n:                                     # bù nếu một nhóm quá nhỏ
        rest = df.drop(out.index)
        out = pd.concat([out, rest.sample(n - len(out), random_state=rng.integers(2**31))])
    return out


def _n1_random(cand: pd.DataFrame, misuse_keys: set, n: int, rng: np.random.Generator) -> pd.DataFrame:
    """N1: địa chỉ ứng viên (cùng nguồn GitHub/StackExchange) có giao dịch nhưng rule tác giả không gắn cờ.
    7,5 triệu dòng → lấy dư một mẫu ngẫu nhiên rồi mới lọc."""
    pool = cand[cand.n_tx > 0].sample(n * 20, random_state=rng.integers(2**31))
    keep = [(c, a) not in misuse_keys for c, a in zip(pool.chain, pool.address)]
    return pool[np.array(keep)].head(n)


def delegations(s: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """Mỗi (chain, EOA) ủy quyền 7702, kèm kết quả symexec của delegate và tên ví hợp lệ đã kiểm chứng (N6).

    `eip7702_delegate_stats.flag` của tác giả chính là kết quả symexec (1 = có đường chuyển tiền, 0 = không,
    -1 = chưa phân tích), không phải nhãn độc: delegate MetaMask hợp lệ cũng có flag 1. Vì vậy dùng symexec làm
    bằng chứng, và tách riêng delegate là ví hợp lệ đã kiểm chứng.
    """
    dl = s["eip7702_delegations"].drop_duplicates(["chain", "eoa"]).rename(columns={"eoa": "address"})
    sx = s["symexec_results"]
    sx = sx[sx.kind == "eip7702_delegate"].drop_duplicates(["chain", "address"])
    transfer = {(c, a): bool(x or y or z) for c, a, x, y, z in
                zip(sx.chain, sx.address, sx.transfer_all_eth, sx.transfer_msg_value, sx.transfer_input_amount)}
    wallets = s["n6"][s["n6"].kept].set_index("address").wallet.to_dict() if len(s["n6"]) else {}
    return dl.assign(delegate_symexec_transfer=[transfer.get((c, d)) for c, d in zip(dl.chain, dl.delegate)],
                     delegate_known_wallet=[wallets.get(d) for d in dl.delegate])


def load_sources() -> dict[str, pd.DataFrame]:
    s = {name: pd.read_parquet(SEED_DIR / f"{name}.parquet") for name in
         ("ca_misuse", "eoa_misuse", "rq1_labels", "symexec_results", "eip7702_delegations",
          "candidate_addresses")}
    s["matrix"] = pd.read_parquet(INTERIM_DIR / "crosschain_matrix.parquet")
    s["n7"] = pd.read_parquet(INTERIM_DIR / "n7_multichain_contracts.parquet")
    n6 = INTERIM_DIR / "n6_benign_7702_delegates.parquet"           # python -m address_misuse.negatives n6
    s["n6"] = pd.read_parquet(n6) if n6.exists() else pd.DataFrame(columns=["address", "wallet", "kept"])
    return s


def draw_pilot(s: dict[str, pd.DataFrame], seed: int = 0) -> pd.DataFrame:
    """Trả về DataFrame (stratum, chain, address, author_label) — 50 cặp (chain, address) không trùng."""
    rng = np.random.default_rng(seed)
    m, rq1 = s["matrix"], s["rq1_labels"]
    eoa = _dedup_max(s["eoa_misuse"])
    dl = delegations(s)
    sweep = (dl.delegate_symexec_transfer == True) & dl.delegate_known_wallet.isna()   # noqa: E712
    misuse_keys = set(zip(s["ca_misuse"].chain, s["ca_misuse"].address)) | set(zip(eoa.chain, eoa.address))
    ncd = m[m.subtype == "NCD"]

    def by_chain(df, n, **kw):
        """chia đều giữa eth và bsc."""
        parts = [_pick(df[df.chain == c], k, rng, **kw) for c, k in (("eth", n - n // 2), ("bsc", n // 2))]
        return pd.concat(parts)

    pools = {
        "rq1_ca_pos": lambda n: _pick(rq1[(rq1.misuse_type == "CA") & (rq1.label == 1)], n, rng, None),
        "rq1_ca_neg": lambda n: _pick(rq1[(rq1.misuse_type == "CA") & (rq1.label == 0)], n, rng, None),
        "rq1_eoa_addr": lambda n: _pick(rq1[rq1.misuse_type == "EOA_ADDR"], n, rng, None),
        "t1_ncd_other": lambda n: by_chain(ncd[ncd.xc_label == "NCD-Other"], n),
        "t4_xc_before": lambda n: _pick(ncd[(ncd.chain == "eth") & (ncd.xc_timing == "deployed_before")], n, rng),
        "t4_xc_bsc": lambda n: _pick(ncd[(ncd.chain == "bsc") & (ncd.xc_label == "NCD-XC")], n, rng),
        "t4_xc_after": lambda n: _pick(ncd[(ncd.chain == "eth") & (ncd.xc_timing == "deployed_after")], n, rng),
        "t4_xc_unknown": lambda n: _pick(ncd[(ncd.chain == "eth") & (ncd.xc_label == "NCD-XC")
                                             & (ncd.xc_timing == "unknown")], n, rng),
        "ca_pcd": lambda n: pd.concat([_pick(m[m.xc_label == "PCD-Other"], n // 2, rng),
                                       _pick(m[m.xc_label == "PCD-XC"], n - n // 2, rng)]),
        "eoa_key": lambda n: by_chain(eoa[eoa.subtype == "KEY"], n),
        "eoa_addr": lambda n: by_chain(eoa[eoa.subtype == "ADDR"], n),
        "eoa_7702": lambda n: _pick(dl[sweep], n, rng, None),
        # ca khó: delegate là ví hợp lệ, không có đường chuyển tiền, hoặc chưa phân tích. Dữ liệu tác giả không có
        # EOA ủy quyền lành tính thật (mọi EOA đều lộ key/ADDR) → N6 đúng nghĩa cần dữ liệu mới (BigQuery, D1).
        "eoa_7702_other": lambda n: _pick(dl[~sweep], n, rng, None),
        "neg_n7_multichain": lambda n: _pick(s["n7"][s["n7"].kept].assign(chain="eth"), n, rng, None),
        "neg_n1_random": lambda n: _n1_random(s["candidate_addresses"], misuse_keys, n, rng),
    }
    rows, seen = [], set()
    for stratum, n in PILOT_PLAN:
        got = pools[stratum](n * 3)                       # dư để còn chỗ bỏ trùng
        k = 0
        for r in got.to_dict("records"):
            key = (r["chain"], r["address"].lower())
            if key in seen:
                continue
            seen.add(key)
            rows.append({"stratum": stratum, "chain": key[0], "address": key[1],
                         "author_label": r.get("label") if stratum.startswith("rq1_") else None})
            k += 1
            if k == n:
                break
        if k < n:
            print(f"  cảnh báo: tầng {stratum} chỉ có {k}/{n} mẫu")
    return pd.DataFrame(rows)


def _state(rec: dict) -> str | None:
    if rec.get("err"):
        return None
    return "7702" if rec.get("d") else ("code" if rec.get("s") else "none")


def live_evidence(chain: str, recs: dict[str, dict]) -> dict:
    """Bằng chứng code của một địa chỉ từ getCode trên mọi mạng (`recs[net]` = `_code_record`).

    Tính cùng một cách cho mọi mẫu (dương, âm, EOA): nếu chỉ một số tầng có ô được điền, người gán đoán ra tầng.
    """
    out = {f"code_{n}": _state(recs[n]) for n in NETWORKS}
    own = recs[chain]
    out["xc_nets"] = ",".join(n for n in NETWORKS if n != chain and out[f"code_{n}"] == "code")
    out["same_code_elsewhere"] = None if own.get("err") else bool(
        own.get("h") and any(recs[n].get("h") == own["h"] for n in NETWORKS if n != chain))
    out["delegate_7702"] = own.get("d")
    return out


def _probe_live(ev: pd.DataFrame) -> dict[tuple[str, str], str]:
    """Ghi đè code_<net>, xc_nets, same_code_elsewhere, delegate_7702 bằng getCode hiện tại cho mọi mẫu.
    Mạng nào RPC lỗi thì giữ trạng thái cũ (ma trận / N7). Trả bytecode đầy đủ trên chain của mẫu (cho symexec)."""
    from address_misuse.crosschain_matrix import _code_record, rpc_batch
    addrs = ev.address.tolist()
    recs, own_code = [{} for _ in addrs], {}
    for n in NETWORKS:
        try:
            codes = rpc_batch(n, [("eth_getCode", [a, "latest"]) for a in addrs])
        except Exception:
            codes = [None] * len(addrs)
        for j, (a, c, ch) in enumerate(zip(addrs, codes, ev.chain)):
            recs[j][n] = _code_record(a, c)
            if n == ch and c:
                own_code[(ch, a)] = c
    for j, (i, ch) in enumerate(zip(ev.index, ev.chain)):
        for col, v in live_evidence(ch, recs[j]).items():
            if col.startswith("code_") and v is None:
                continue                                   # RPC lỗi: giữ trạng thái từ ma trận nếu có
            ev.at[i, col] = v
    return own_code


SYMEXEC_CACHE = RAW_DIR / "gold" / "symexec_pilot.jsonl"


def _symexec_missing(tasks: list[tuple[str, str, str]], timeout: float = 120, workers: int = 2) -> dict:
    """Chạy symexec (engine của tác giả) cho các (chain, address, bytecode) chưa có kết quả của tác giả.
    Có cache: chạy lại `pilot` không phân tích lại."""
    from address_misuse.reproduce_symexec import run_isolated
    done = {}
    if SYMEXEC_CACHE.exists():
        for line in SYMEXEC_CACHE.read_text(encoding="utf-8").splitlines():
            r = json.loads(line)
            done[(r["chain"], r["address"])] = r
    todo = [(c, "pilot", a, code, timeout) for c, a, code in tasks if (c, a) not in done]
    if todo:
        print(f"  symexec {len(todo)} contract chưa có kết quả của tác giả (≤ {timeout + 60:.0f}s mỗi contract)…",
              flush=True)
        SYMEXEC_CACHE.parent.mkdir(parents=True, exist_ok=True)
        with SYMEXEC_CACHE.open("a", encoding="utf-8") as f:
            for r in run_isolated(todo, workers=workers, hard_limit=timeout + 60, max_mem_mb=1500):
                done[(r["chain"], r["address"])] = r
                f.write(json.dumps(r) + "\n")
                f.flush()
    return {k: {f: r[f"ours_{f}"] for f in SYMEXEC_FLAGS} for k, r in done.items() if r.get("status") == "ok"}


def _author_symexec(s: dict[str, pd.DataFrame]) -> dict[tuple[str, str], dict]:
    sx = s["symexec_results"].drop_duplicates(["chain", "address"])
    return {(c, a): {f: bool(v) for f, v in zip(SYMEXEC_FLAGS, vals)}
            for c, a, *vals in zip(sx.chain, sx.address, *(sx[f] for f in SYMEXEC_FLAGS))}


def _transfer(flags: dict | None) -> bool | None:
    return None if flags is None else bool(flags["transfer_all_eth"] or flags["transfer_msg_value"]
                                           or flags["transfer_input_amount"])


XC_ORIGIN_COLS = ["deployer", "create_kind", "origin_net", "first_create_ts", "first_tx_ts", "xc_timing", "redeploy"]


def add_evidence(sample: pd.DataFrame, s: dict[str, pd.DataFrame], run_symexec: bool = True,
                 fallback: bool = True) -> pd.DataFrame:
    """`fallback=False` (mẫu chính): KHÔNG điền sẵn trạng thái code từ ma trận / N7 khi RPC lỗi, vì hai nguồn đó chỉ có cho một số
    tầng (ma trận = địa chỉ misuse của tác giả) nên ô được điền hay không sẽ làm lộ tầng. Lỗi RPC để None, nơi gọi phải xử lý."""
    keys = ["chain", "address"]
    m = s["matrix"].copy()
    ev = sample[keys].copy()

    # Không đưa n_tx / n_misuse_tx / cost_native của tác giả vào file mẫu: chúng chỉ có cho địa chỉ trong seed,
    # nên ô trống ở mẫu âm và số "misuse tx" > 0 ở mẫu dương đều làm lộ nhãn.

    code_cols = {f"{n}_state": f"code_{n}" for n in NETWORKS}
    m = m.reindex(columns=keys + list(code_cols) + ["deployer", "kind", "origin_net", "first_create_ts",
                                                   "first_tx_ts", "xc_timing", "redeploy"])
    ev = ev.merge(m.rename(columns={**code_cols, "kind": "create_kind"}), on=keys, how="left")
    ev = ev.reindex(columns=keys + list(code_cols.values()) + ["xc_nets", "same_code_elsewhere", "delegate_7702"]
                    + XC_ORIGIN_COLS).astype(object)

    if fallback:
        n7 = s["n7"].set_index("address")
        for i, r in ev.iterrows():                       # dự phòng khi RPC lỗi: N7 không nằm trong ma trận misuse
            if r.address in n7.index and pd.isna(r.code_eth):
                nets = n7.loc[r.address, "networks_with_code"].split(",")
                for n in NETWORKS:
                    ev.at[i, f"code_{n}"] = "code" if n in nets else "none"
    else:
        ev[list(code_cols.values())] = None
    own_code = _probe_live(ev)
    own_state = [ev.at[i, f"code_{c}"] for i, c in zip(ev.index, ev.chain)]

    # Nguồn gốc / thời điểm deploy ở chain khác chỉ áp dụng khi chain này KHÔNG có code mà chain khác có
    # (tình huống T4). Ngoài tình huống đó để trống cho mọi mẫu, kể cả mẫu trong ma trận.
    xc = (ev.xc_nets.fillna("") != "") & (pd.Series(own_state, index=ev.index) != "code")
    ev.loc[~xc, XC_ORIGIN_COLS] = None

    # symexec cho contract trên chain này và cho delegate 7702: ưu tiên kết quả của tác giả, thiếu thì tự chạy.
    author = _author_symexec(s)
    need = {(c, a) for c, a, st in zip(ev.chain, ev.address, own_state) if st == "code" and (c, a) not in author}
    need_d = {(c, d) for c, d in zip(ev.chain, ev.delegate_7702) if isinstance(d, str) and (c, d) not in author}
    if need_d and run_symexec:
        from address_misuse.crosschain_matrix import rpc_batch
        for c in sorted({c for c, _ in need_d}):
            ds = sorted(d for cc, d in need_d if cc == c)
            for d, code in zip(ds, rpc_batch(c, [("eth_getCode", [d, "latest"]) for d in ds])):
                if code and code != "0x":
                    own_code[(c, d)] = code
    tasks = [(c, a, own_code[(c, a)]) for c, a in sorted(need | need_d) if (c, a) in own_code]
    flags = {**(_symexec_missing(tasks) if run_symexec and tasks else {}), **author}
    for f in SYMEXEC_FLAGS:
        ev[f"symexec_{f}"] = [flags.get((c, a), {}).get(f) if st == "code" else None
                              for c, a, st in zip(ev.chain, ev.address, own_state)]
    wallets = s["n6"][s["n6"].kept].set_index("address").wallet.to_dict() if len(s["n6"]) else {}
    ev["delegate_symexec_transfer"] = [_transfer(flags.get((c, d))) if isinstance(d, str) else None
                                       for c, d in zip(ev.chain, ev.delegate_7702)]
    ev["delegate_known_wallet"] = [wallets.get(d) if isinstance(d, str) else None for d in ev.delegate_7702]

    # Không có cột "nằm trong danh sách khoá lộ": danh sách chỉ chứa dương KEY nên cột đó làm lộ nhãn (FINDINGS_LOG 2026-10-25).

    ev["first_create_date"] = ev.first_create_ts.map(_date)
    ev["first_tx_date"] = ev.first_tx_ts.map(_date)
    ev = ev.drop(columns=["first_create_ts", "first_tx_ts"])
    ev["explorer_url"] = [EXPLORERS[c] + a for c, a in zip(ev.chain, ev.address)]
    ev["explorer_url_xc"] = [EXPLORERS[o] + a if isinstance(o, str) and o in EXPLORERS else None
                             for o, a in zip(ev.origin_net, ev.address)]
    return ev


def cmd_pilot(seed: int = 0) -> pd.DataFrame:
    s = load_sources()
    sample = draw_pilot(s, seed)
    ev = add_evidence(sample, s)
    order = np.random.default_rng(seed + 1).permutation(len(sample))   # xáo để người gán không đoán được tầng
    sample = sample.iloc[order].reset_index(drop=True)
    sample.insert(0, "item_id", [f"P{i + 1:03d}" for i in range(len(sample))])
    items = sample[["item_id", "chain", "address"]].merge(ev, on=["chain", "address"], how="left")
    items = items.astype(object).where(items.notna() & (items != ""), "-")      # "-" = không tra được, không suy ra gì từ ô này
    GOLD_DIR.mkdir(parents=True, exist_ok=True)
    items.to_csv(GOLD_DIR / "pilot50_items.csv", index=False, encoding="utf-8")
    sample[["item_id", "chain", "address", "stratum", "author_label"]].to_csv(
        GOLD_DIR / "pilot50_key.csv", index=False, encoding="utf-8")
    pd.DataFrame({"item_id": items.item_id, **{c: "" for c in LABEL_COLS}}).to_csv(
        GOLD_DIR / "annotator_template.csv", index=False, encoding="utf-8")
    print(f"Ghi {len(items)} mẫu → {GOLD_DIR / 'pilot50_items.csv'} (+ pilot50_key.csv, annotator_template.csv)")
    print(sample.stratum.value_counts().sort_index().to_string())
    return items


# ---------------------------------------------------------------- mẫu vàng chính (từ dataset_v2)

# Kế hoạch mẫu (tiền đăng ký, GOLD_PROTOCOL §1). Đơn vị lấy mẫu là địa chỉ (chain, address); mỗi địa chỉ thuộc đúng một tầng
# (tầng đầu tiên trong kế hoạch có địa chỉ đó). Tầng "+D6" = dương mới D6 (label_source == d6_rule).
# Khung = test ∪ out_of_time của dataset_v2 (bỏ excluded): nhãn vàng dùng để chấm engine nên phải ở tập đánh giá.
# - Dương (mọi tầng "+"): lấy HẾT (điều tra toàn bộ, trọng số 1).
# - Âm bị engine gắn cờ (luật cứng hoặc R ≥ τ_warn): lấy HẾT.
# - Âm không cờ: mẫu ngẫu nhiên đơn giản trong từng ô (tầng × chain), số địa chỉ ở lô 1 là giá trị dưới đây
#   (None = lấy hết); phần còn lại xếp thành danh sách dự phòng đã xáo cho lô 2 (`draw_reserve`).
MAIN_PLAN = {"+ADDR": None, "+KEY": None, "+NCD": None, "+PCD": None, "+D6": None, "-N1": 40, "-N2": 60, "-N3": 40,
             "-N4": 30, "-N5": None, "-N6": None, "-N7": None, "-N8": 80}
RESERVE_BATCH = 50                                        # lô 2 gán theo từng lô 50 mẫu, chỉ dừng ở ranh giới lô


def _take_counts(sizes: dict[str, int], n: int) -> dict[str, int]:
    """Chia n địa chỉ đều giữa eth và bsc; chain nào thiếu thì chain kia bù."""
    want = {"eth": n - n // 2, "bsc": n // 2}
    k = {c: min(want[c], sizes.get(c, 0)) for c in want}
    short = n - sum(k.values())
    for c in want:
        extra = min(short, sizes.get(c, 0) - k[c])
        k[c] += extra
        short -= extra
    return k


def _units(ds: pd.DataFrame, seed: int = 0) -> pd.DataFrame:
    """Khung lấy mẫu theo ĐỊA CHỈ: mỗi dòng là một (chain, address) với dòng đại diện (ưu tiên dòng engine gắn cờ), số dòng
    của địa chỉ trong khung (`n_rows`), ô lấy mẫu (`cell` = tầng|chain|F/u), thứ tự xáo trong ô (`rank`), cỡ ô (`cell_n`,
    `cell_rows`) và số địa chỉ lô 1 lấy trong ô (`k`). Lô 1 = rank < k; danh sách dự phòng = âm không cờ có rank ≥ k."""
    rng = np.random.default_rng(seed)
    pool = ds[(ds.split == "test") | ds.out_of_time].copy()
    pool = pool[pool.split != "excluded"]
    pool["stratum"] = np.where(pool.label_source == "d6_rule", "+D6", pool.grp)
    pool["flag"] = pool["flag"].fillna(False).astype(bool) if "flag" in pool else False
    pool = pool.sample(frac=1, random_state=int(rng.integers(2**31)))        # xáo trước khi bỏ trùng
    pool = pool.sort_values("flag", ascending=False, kind="stable")           # dòng có cờ lên trước → làm dòng đại diện
    # số dòng của địa chỉ trong TOÀN khung (không chỉ trong tầng của nó): Σ n_rows = cỡ khung; nhãn là của địa chỉ
    rows_all = pool.groupby(["chain", "address"]).size().rename("n_rows").reset_index()
    taken, parts = set(), []
    for stratum, n in MAIN_PLAN.items():
        d = pool[pool.stratum == stratum]
        d = d[np.array([k not in taken for k in zip(d.chain, d.address)], dtype=bool)]   # np: list rỗng sẽ chọn 0 cột
        u = d.drop_duplicates(["chain", "address"]).merge(rows_all, on=["chain", "address"], how="left")
        taken |= set(zip(u.chain, u.address))
        u["cell"] = stratum + "|" + u.chain + "|" + np.where(u.flag, "F", "u")
        u["rank"] = u.groupby("cell").cumcount()
        u["cell_n"] = u.groupby("cell").cell.transform("size")
        u["cell_rows"] = u.groupby("cell").n_rows.transform("sum")
        u["k"] = u.cell_n                                              # mặc định: lấy hết ô
        if stratum[0] == "-" and n is not None:                        # âm không cờ: lấy mẫu
            sizes = u[~u.flag].groupby("chain").size().to_dict()
            kk = _take_counts(sizes, n)
            if len(u) and sum(kk.values()) < n:
                print(f"  cảnh báo: tầng {stratum} chỉ có {sum(kk.values())}/{n} địa chỉ không cờ")
            m = ~u.flag
            u.loc[m, "k"] = u.loc[m, "chain"].map(kk).fillna(0).astype(int)
        parts.append(u)
    return pd.concat(parts, ignore_index=True)


def draw_main(ds: pd.DataFrame, seed: int = 0, n_double: int = 200) -> pd.DataFrame:
    """Lô 1 của mẫu vàng chính. `w_addr` = cỡ ô / số địa chỉ lấy (ước lượng theo địa chỉ); `w_row` = w_addr × số dòng của
    địa chỉ trong khung (ước lượng theo dòng, như các chỉ số của engine). Dương và âm có cờ: trọng số 1 theo địa chỉ.
    `n_double` mẫu gán đôi rút ngẫu nhiên từ cả lô. Thứ tự đã xáo (dương và âm lẫn lộn, giữ làm mù)."""
    u = _units(ds, seed)
    out = u[u["rank"] < u.k].copy()
    out["w_addr"] = out.cell_n / out.k
    out["w_row"] = out.w_addr * out.n_rows
    out["weight"] = out.w_row                                   # tên cũ
    out = out.iloc[np.random.default_rng(seed + 1).permutation(len(out))].reset_index(drop=True)
    out["double"] = False
    out.loc[out.sample(min(n_double, len(out)), random_state=seed + 2).index, "double"] = True
    # Danh sách bù (rút sẵn, thứ tự cố định): nếu cổng κ ở 50 mẫu đầu bị dừng thì 50 mẫu gán đôi đã dùng để sửa codebook bị bỏ,
    # và thay bằng các mẫu đơn có `double_extra_rank` nhỏ nhất (GOLD_PROTOCOL §2.4). -1 = đã là mẫu gán đôi.
    rest = out.index[~out.double.to_numpy()]
    out["double_extra_rank"] = -1
    out.loc[rest[np.random.default_rng(seed + 4).permutation(len(rest))], "double_extra_rank"] = np.arange(len(rest))
    out.insert(0, "item_id", [f"G{i + 1:03d}" for i in range(len(out))])
    return out


def draw_reserve(ds: pd.DataFrame, seed: int = 0, double_frac: float = 0.1) -> pd.DataFrame:
    """Danh sách âm dự phòng cho lô 2, thứ tự cố định tính từ trước khi gán: trong từng ô là tiếp nối thứ tự xáo của lô 1
    (nên tập đã gán của ô luôn là mẫu ngẫu nhiên đơn giản), giữa các ô xen kẽ theo tỉ lệ cỡ ô (mọi tiền tố ≈ phân bổ tỉ lệ).
    Chia lô `RESERVE_BATCH`; chỉ dừng ở ranh giới lô, theo ngân sách / hạn chót, KHÔNG theo kết quả engine."""
    u = _units(ds, seed)
    r = u[(u.stratum.str[0] == "-") & ~u.flag & (u["rank"] >= u.k)].copy()
    r["pos"] = (r["rank"] - r.k + 0.5) / (r.cell_n - r.k)
    r = r.sort_values("pos", kind="stable").reset_index(drop=True)
    r["batch"] = np.arange(len(r)) // RESERVE_BATCH + 1
    r["double"] = np.random.default_rng(seed + 3).random(len(r)) < double_frac
    r.insert(0, "item_id", [f"R{i + 1:04d}" for i in range(len(r))])
    return r


def engine_flag(ds: pd.DataFrame) -> np.ndarray:
    """Cờ cảnh báo của engine (mức warn: luật cứng hoặc R ≥ τ_warn) cho từng dòng dataset_v2. `engine_eval.parquet` cùng thứ tự."""
    ev = pd.read_parquet(INTERIM_DIR / "engine_eval.parquet")
    for c in ("chain", "address", "grp", "at"):
        if not (ds[c].values == ev[c].values).all():
            raise SystemExit(f"engine_eval.parquet lệch dataset_v2 ở cột {c}: chạy lại engine.evaluate")
    tau = json.load(open(INTERIM_DIR / "engine_metrics.json", encoding="utf-8"))["main"]["thresholds"]["warn"]
    return ev.hard.notna().to_numpy() | (ev.R.to_numpy() >= tau)


def event_tx(chain: str, address: str, at) -> dict | None:
    """Tx đi vào `address` tại đúng thời điểm `at` (mẫu theo địa chỉ chọn `at` từ một tx đi vào, nhưng dataset không lưu hash).
    Cần cho người gán thấy cùng loại bằng chứng ở mọi mẫu: nếu chỉ mẫu âm theo tx có tx_hash thì có hash = âm."""
    import gzip
    from address_misuse.engine.features import txlist_path
    p = txlist_path(chain, address)
    if not p.exists() or at is None or pd.isna(at):
        return None
    with gzip.open(p, "rt", encoding="utf-8") as f:
        txs = json.load(f)["result"]
    hit = [t for t in txs if (t.get("to") or "").lower() == address and int(t["timeStamp"]) == int(at)]
    hit.sort(key=lambda t: (int(t["gasUsed"]) <= 21_000, int(t.get("transactionIndex") or 0)))   # ưu tiên tx có calldata
    return hit[0] if hit else None


SYMEXEC_COLS = [f"symexec_{f}" for f in SYMEXEC_FLAGS]
LIVE_COLS = [f"code_{n}" for n in NETWORKS] + ["xc_nets", "same_code_elsewhere", "delegate_7702"] + SYMEXEC_COLS
# Cột của main_items.csv (v0.6), đúng thứ tự. Mọi cột theo thời điểm tính bằng `gold_evidence` cho MỌI mẫu theo cùng quy tắc
# (GOLD_PROTOCOL §3); cột có điều kiện chỉ phụ thuộc vào các cột khác người gán cũng thấy, không phụ thuộc tầng.
MAIN_COLS = (["item_id", "chain", "address", "at_time", "tx_hash", "tx_url", "tx_status", "selector", "calldata_kind",
              "calldata_len", "value_native", "code_at_event", "own_create_time", "creator"] + [f"code_{n}" for n in NETWORKS]
             + ["xc_nets", "same_code_elsewhere", "delegate_7702", "xc_timing", "xc_create_time", "origin_net", "deployer"]
             + SYMEXEC_COLS + ["explorer_url", "explorer_url_xc"])
XC_COLS = ["xc_timing", "xc_create_time", "origin_net", "deployer", "explorer_url_xc"]


def main_evidence(sample: pd.DataFrame, s: dict[str, pd.DataFrame], live_from: str | None = None) -> pd.DataFrame:
    """Bằng chứng cho mẫu chính (CODEBOOK v0.6 §6, GOLD_PROTOCOL §3). Không đưa nhãn, tầng, điểm R hay số liệu của tác giả.

    - Theo địa chỉ, trạng thái HÔM NAY (`LIVE_COLS`): getCode 8 mạng + symexec (`add_evidence`, không dùng nguồn dự phòng chỉ có ở
      một số tầng). `live_from` = một main_items.csv cũ: dùng lại đúng các cột này (cùng một ảnh chụp cho mọi mẫu) thay vì tra lại,
      để bảng tra delegate không đổi giữa hai phiên bản.
    - Theo tx tại `at` (`gold_evidence`): giờ tx, trạng thái, selector, loại calldata, value; trạng thái code NGAY TRƯỚC tx
      (`code_at_event`; với địa chỉ hôm nay là 7702, so thêm lúc tạo contract delegate) và lúc tạo contract trên chain này
      (`own_create_time`, `creator`) tra cho MỌI mẫu;
      contract cùng địa chỉ ở mạng khác so với tx tại `at` (`xc_*`) cho MỌI mẫu chain này không có code lúc tx mà mạng khác có code.
    """
    from address_misuse import gold_evidence as ge
    keys = ["chain", "address"]
    addrs = sample[keys].drop_duplicates()
    if live_from:
        old = pd.read_csv(live_from, dtype=str).fillna("-")
        live = addrs.merge(old[keys + LIVE_COLS].drop_duplicates(keys), on=keys, how="left")
        if live[LIVE_COLS].isna().any().any():
            raise SystemExit(f"{live_from} thiếu cột trạng thái hôm nay cho một số địa chỉ: tra lại (bỏ --live-from)")
    else:
        live = add_evidence(addrs, s, fallback=False)[keys + LIVE_COLS]
        bad = live[[f"code_{n}" for n in NETWORKS]].isna().any(axis=1)
        if bad.any():                                    # không để ô trống không đều giữa các tầng: chạy lại khi RPC ổn định
            raise SystemExit(f"getCode lỗi cho {int(bad.sum())} địa chỉ: chạy lại `sample_gold main`")
    live = live.astype(object).where(live.notna() & (live != ""), "-")
    sample = sample.merge(live, on=keys, how="left")
    # dòng của dataset theo địa chỉ không lưu hash: khôi phục tx đi vào tại đúng `at` từ txlist (mọi mẫu đều phải có tx_hash)
    sample["tx_hash"] = [h if isinstance(h, str) else ((event_tx(c, a, at) or {}).get("hash"))
                         for h, c, a, at in zip(sample.tx_hash, sample.chain, sample.address, sample["at"])]

    txs = ge.fetch_event_txs([(c, h) for c, h in zip(sample.chain, sample.tx_hash) if isinstance(h, str)])
    own = ge.fetch_creations(list(zip(sample.chain, sample.address)))
    # v0.6.2: bản ghi tạo contract của delegate 7702 (cùng chain), để biết delegate đã tồn tại lúc tx chưa
    dpairs = sorted({(c, str(d).lower()) for c, d in zip(sample.chain, sample.delegate_7702) if str(d) not in ("-", "nan", "")})
    dele = ge.fetch_creations(dpairs) if dpairs else {}
    rows = []
    for r in sample.to_dict("records"):
        c, a, h, at = r["chain"], r["address"], r["tx_hash"], int(r["at"])
        t = txs.get((c, h)) if isinstance(h, str) else None
        cr = own.get((c, a))
        if cr and cr.get("status") == "found" and t and cr.get("block") == t["block"]:
            cr = ge.creation_index(c, cr)                # cùng block: cần vị trí trong block để biết cái nào trước
        dc = dele.get((c, str(r["delegate_7702"]).lower()))
        if dc and dc.get("status") == "found" and t and dc.get("block") == t["block"]:
            dc = ge.creation_index(c, dc)
        cae = ge.state_at_event(c, r[f"code_{c}"], at, t, cr, dc)
        found = bool(cr) and cr.get("status") == "found"
        inp = t["input"] if t else None
        row = {"item_id": r["item_id"], "at_time": ge.utc(at), "tx_hash": h if isinstance(h, str) else None,
               "tx_url": EXPLORERS[c].replace("/address/", "/tx/") + h if isinstance(h, str) else None,
               "tx_status": None if not t else ("success" if t["status"] == 1 else "fail"),
               "selector": inp[:10] if inp and len(inp) >= 10 else None, "calldata_kind": ge.calldata_kind(inp),
               "calldata_len": ge.calldata_len(inp),
               "value_native": int(t["value"]) / 1e18 if t else None, "code_at_event": cae,
               "own_create_time": ge.utc(cr["ts"]) if found else None, "creator": cr.get("creator") if found else None}
        nets = [n for n in str(r["xc_nets"]).split(",") if n and n != "-"]
        if nets and cae not in ("code", "7702"):         # điều kiện chỉ dựa vào hai cột người gán cũng thấy
            xc = ge.fetch_creations([(n, a) for n in nets])
            f = ge.xc_fields(at, nets, {n: xc.get((n, a)) for n in nets})
            link = f.pop("xc_net_link")
            row.update(f, explorer_url_xc=EXPLORERS[link] + a if link in EXPLORERS else None)
        rows.append(row)
    items = sample[["item_id", "chain", "address"]].merge(pd.DataFrame(rows), on="item_id").merge(live, on=keys, how="left")
    items["explorer_url"] = [EXPLORERS[c] + a for c, a in zip(items.chain, items.address)]
    items = items.reindex(columns=MAIN_COLS)
    items["calldata_len"] = items.calldata_len.map(lambda v: None if pd.isna(v) else str(int(v)))   # không thành "64.0"
    return items.astype(object).where(items.notna() & (items != ""), "-")


DELEGATE_REVIEW_COLS = {"A": ["class_A", "note_A"], "B": ["class_B", "note_B"]}
DELEGATE_KEEP = ["class_A", "note_A", "class_B", "note_B", "final_class", "source", "note"]
DELEGATE_FIXED = ["chain", "delegate", "n_items", "explorer_url", "final_class", "source", "note"]


def write_delegate_lookup(items: pd.DataFrame) -> pd.DataFrame:
    """Bảng tra delegate EIP-7702 (CODEBOOK ca khó 13, GOLD_PROTOCOL §2.3): mỗi (chain, delegate) trong mẫu chính một dòng.
    Người điều phối điền sẵn `final_class`, `source` và `note` (lý do, từ mã nguồn đã verify hoặc bytecode). A và B RÀ SOÁT
    trước khi gán: mỗi dòng ghi `class_A` / `class_B` = `ok` (đã xem, đồng ý) hoặc một lớp khác kèm lý do ở cột của chính mình
    (`note_A` / `note_B`) = phản đối. Ô trống = CHƯA RÀ (không phải đồng ý). Lớp: legit_wallet / forwarder_sweeper / unknown.
    Chạy lại GIỮ NGUYÊN mọi ô đã điền (kể cả `note_A`, `note_B`); delegate mới thêm dòng trống. Không chứa nhãn mẫu."""
    d = items[items.delegate_7702.notna() & (items.delegate_7702 != "-")]
    t = (d.groupby(["chain", "delegate_7702"]).agg(n_items=("item_id", "size"))
         .reset_index().rename(columns={"delegate_7702": "delegate"}).sort_values("n_items", ascending=False, kind="stable"))
    t["explorer_url"] = [EXPLORERS[c] + a for c, a in zip(t.chain, t.delegate)]
    p = GOLD_DIR / "delegate_lookup.csv"
    if p.exists():
        old = pd.read_csv(p, dtype=str, keep_default_na=False).reindex(columns=["chain", "delegate", *DELEGATE_KEEP])
        old["_pos"] = range(len(old))                    # giữ thứ tự dòng cũ (chỉ khác ở các dòng bằng n_items)
        t = t.merge(old[["chain", "delegate", "_pos", *DELEGATE_KEEP]], on=["chain", "delegate"], how="left")
        t = t.sort_values(["_pos", "n_items"], ascending=[True, False], na_position="last", kind="stable")
        gone = set(zip(old.chain, old.delegate)) - set(zip(t.chain, t.delegate))
        if gone:
            print(f"  cảnh báo: {len(gone)} delegate cũ không còn trong mẫu (bỏ khỏi bảng): {sorted(gone)[:3]}")
    t = t.reindex(columns=["chain", "delegate", "n_items", "explorer_url", *DELEGATE_KEEP]).fillna("")
    new = int((t.final_class == "").sum())
    if new:
        print(f"  {new} delegate chưa có final_class: người điều phối phải phân loại trước khi giao Excel")
    t.to_csv(p, index=False, encoding="utf-8", lineterminator="\n")
    return t


DELEGATE_CLASSES = ("legit_wallet", "forwarder_sweeper", "unknown")


def check_delegates(t: pd.DataFrame, who: tuple[str, ...] = ("A", "B"), strict: bool = False) -> list[str]:
    """Bảng delegate đã được rà chưa (CODEBOOK ca 13): với mỗi người trong `who`, mỗi dòng `class_X` phải là `ok` (đồng ý) hoặc
    một lớp khác `final_class` kèm lý do ở `note_X` của chính người đó (phản đối). Ô trống = chưa rà. `final_class` phải hợp lệ.
    `strict` (cổng trước khi tạo file gán, GOLD_PROTOCOL §2.3): mọi phản đối phải đã giải quyết, tức `class_X` ∈ {ok, final_class};
    còn phản đối nào (kể cả có lý do) là lỗi."""
    errs = []
    t = t.reindex(columns=list(dict.fromkeys([*t.columns, "final_class", *(c for w in who for c in DELEGATE_REVIEW_COLS[w])])))
    for r in t.fillna("").itertuples(index=False):
        g = r._asdict()
        if g["final_class"] not in DELEGATE_CLASSES:
            errs.append(f"{r.chain} {r.delegate}: final_class = {g['final_class']!r} không hợp lệ")
        for w in who:
            c, n = DELEGATE_REVIEW_COLS[w]
            v = str(g[c]).strip().lower()
            if not v:
                errs.append(f"{r.chain} {r.delegate}: {c} trống (chưa rà)")
            elif v != "ok" and v not in DELEGATE_CLASSES:
                errs.append(f"{r.chain} {r.delegate}: {c} = {g[c]!r} (chỉ nhận ok / {' / '.join(DELEGATE_CLASSES)})")
            elif v not in ("ok", g["final_class"]) and not str(g[n]).strip():
                errs.append(f"{r.chain} {r.delegate}: {c} phản đối ({v}) nhưng chưa có lý do ở {n}")
            elif strict and v not in ("ok", g["final_class"]):
                errs.append(f"{r.chain} {r.delegate}: {c} phản đối ({v}) chưa giải quyết (cần thống nhất: người điều phối "
                            f"sửa final_class hoặc người rà ghi lại ok)")
    return errs


def cmd_check_delegates(path: str | None = None, who: str | None = None, exit_on_error: bool = True,
                        strict: bool = False) -> list[str]:
    """In lỗi rà soát; có lỗi thì thoát mã 1 (để chặn quy trình). `who` = A hoặc B khi kiểm bản sao riêng của một người.
    `strict`: thêm lỗi cho phản đối chưa giải quyết (đúng cổng mà `sheet --set main` dùng)."""
    t = pd.read_csv(path or GOLD_DIR / "delegate_lookup.csv", dtype=str, keep_default_na=False)
    errs = check_delegates(t, (who,) if who else ("A", "B"), strict=strict)
    print("Bảng delegate đã rà đủ." if not errs else f"{len(errs)} lỗi:\n  " + "\n  ".join(errs))
    if errs and exit_on_error:
        raise SystemExit(1)
    return errs


def cmd_delegate_copy(who: str, out: str | None = None) -> str:
    """Bản sao riêng để `who` rà (GOLD_PROTOCOL §2.3): chỉ có cột của người đó, không thấy cột của người kia."""
    from pathlib import Path
    full = pd.read_csv(GOLD_DIR / "delegate_lookup.csv", dtype=str, keep_default_na=False)
    t = full.reindex(columns=[*DELEGATE_FIXED, *DELEGATE_REVIEW_COLS[who]]).fillna("")
    p = out or str(GOLD_DIR / f"delegate_review_{who}.csv")
    if Path(p).exists():
        raise SystemExit(f"{p} đã có, không ghi đè (có thể đã có ý kiến rà soát)")
    t.to_csv(p, index=False, encoding="utf-8", lineterminator="\n")
    print(f"Ghi bản rà của {who} → {p}")
    return p


def cmd_merge_delegates(path_a: str, path_b: str) -> pd.DataFrame:
    """Gộp hai bản rà riêng vào delegate_lookup.csv: lấy class_A/note_A từ bản của A, class_B/note_B từ bản của B.
    Từ chối nếu bản sao đổi dòng hoặc đổi cột của người điều phối (`final_class`, `source`, `note`)."""
    p = GOLD_DIR / "delegate_lookup.csv"
    m = pd.read_csv(p, dtype=str, keep_default_na=False).reindex(columns=["chain", "delegate", "n_items", "explorer_url",
                                                                          *DELEGATE_KEEP]).fillna("")
    for who, path in (("A", path_a), ("B", path_b)):
        c = pd.read_csv(path, dtype=str, keep_default_na=False)
        miss = set(["chain", "delegate", "final_class", "source", "note", *DELEGATE_REVIEW_COLS[who]]) - set(c.columns)
        if miss:
            raise SystemExit(f"{path}: thiếu cột {sorted(miss)}")
        j = m[["chain", "delegate", "final_class", "source", "note"]].merge(
            c, on=["chain", "delegate"], how="outer", suffixes=("", "_copy"), indicator=True)
        if (j._merge != "both").any():
            raise SystemExit(f"{path}: tập delegate khác bảng gốc")
        for col in ("final_class", "source", "note"):
            if (j[col] != j[f"{col}_copy"]).any():
                raise SystemExit(f"{path}: cột {col} khác bảng gốc (người rà không được sửa cột của người điều phối)")
        got = c.set_index(["chain", "delegate"])[DELEGATE_REVIEW_COLS[who]]
        idx = pd.MultiIndex.from_frame(m[["chain", "delegate"]])
        for col in DELEGATE_REVIEW_COLS[who]:
            m[col] = got[col].reindex(idx).to_numpy()
    m.to_csv(p, index=False, encoding="utf-8", lineterminator="\n")
    errs = check_delegates(m)
    print(f"Gộp xong → {p}; " + ("đã rà đủ." if not errs else f"còn {len(errs)} lỗi (chạy check-delegates để xem)."))
    return m


def cmd_main(seed: int = 0, n_double: int = 200, live_from: str | None = None) -> pd.DataFrame:
    import hashlib
    from address_misuse.config import FINAL_DIR
    p = FINAL_DIR / "dataset_v2.parquet"
    if not p.exists():
        raise SystemExit("chưa có dataset_v2.parquet: chạy address_misuse.freeze (sau khi engine.evaluate xong)")
    ds = pd.read_parquet(p)
    ds["flag"] = engine_flag(ds)
    sample = draw_main(ds, seed, n_double)
    reserve = draw_reserve(ds, seed)
    items = main_evidence(sample, None if live_from else load_sources(), live_from)
    GOLD_DIR.mkdir(parents=True, exist_ok=True)
    write_delegate_lookup(items)
    # Từ v0.5 không có `redeploy`, `delegate_symexec_transfer`, `delegate_known_wallet`; từ v0.6 không có `create_kind`,
    # `first_tx_date` (T4 so với tx tại `at`, không so với tx đầu tiên) và các cột ngày được thay bằng giờ UTC (GOLD_PROTOCOL §3).
    items.to_csv(GOLD_DIR / "main_items.csv", index=False, encoding="utf-8")
    keycols = ["item_id", "sample_id", "stratum", "grp", "chain", "address", "label", "label_source", "split",
               "out_of_time", "needs_review", "flag", "cell", "cell_n", "cell_rows", "k", "n_rows", "w_addr", "w_row",
               "double", "double_extra_rank"]
    sample[keycols].to_csv(GOLD_DIR / "main_key.csv", index=False, encoding="utf-8")
    reserve[keycols[:-4] + ["batch", "double"]].to_csv(GOLD_DIR / "main_reserve_key.csv", index=False, encoding="utf-8")
    print(f"Ghi {len(items)} mẫu lô 1 ({int(sample.double.sum())} gán đôi để tính κ) → {GOLD_DIR / 'main_items.csv'} "
          f"(+ main_key.csv); dự phòng lô 2: {len(reserve)} âm ({reserve.batch.max()} lô × {RESERVE_BATCH}) → main_reserve_key.csv")
    for f in ("main_items.csv", "main_key.csv", "main_reserve_key.csv"):
        print(f"  sha256 {f}: {hashlib.sha256((GOLD_DIR / f).read_bytes()).hexdigest()}")
    print(sample.groupby("stratum").agg(n=("item_id", "size"), flag=("flag", "sum"), rows=("n_rows", "sum"),
                                        w_addr_max=("w_addr", "max")).round(2).to_string())
    return items


# ---------------------------------------------------------------- file Excel cho người gán

CHOICES = {"is_misuse": ["1", "0", "?"], "type": ["T1", "T2", "T3", "T4", "T5", "T6", "NONE"],
           "subtype": ["NCD", "PCD", "KEY", "ADDR", "7702", "-"], "malicious": ["1", "0", "?", "-"],
           "confidence": ["high", "medium", "low"]}
SHEET = "Gan_nhan"
GUIDE = [
    "HƯỚNG DẪN GÁN NHÃN (tóm tắt docs/CODEBOOK.md v0.6.2; đọc bản đầy đủ nếu phân vân)",
    "",
    "Gán ĐỘC LẬP. Không mở *_key.csv, file của người gán kia, dataset_v2.parquet, engine_eval.parquet; không đọc GOLD_PROTOCOL.md, FINDINGS_LOG.md, PROGRESS.md trước khi xong.",
    "Chỉ điền các cột màu vàng ở sheet Gan_nhan. Các cột còn lại là bằng chứng, tính cùng một cách cho mọi mẫu; ô '-' = không áp dụng / không tra được, không suy ra gì.",
    "code_<chain> là trạng thái HÔM NAY. Trạng thái lúc tx trên chain này: code_at_event. Chain khác so với tx: xc_timing.",
    "",
    "is_misuse: 1 / 0 / ?   (? chỉ khi không mở được explorer; tối đa khoảng 10% số mẫu của bạn)",
    "type: một type chính T1…T6, hoặc NONE khi is_misuse = 0",
    "  T1 gọi hàm (calldata_kind = call) tới địa chỉ không có code ngay trước tx; không có contract ở chain khác tạo trước tx",
    "  T2 tx THÀNH CÔNG tới địa chỉ có code mà người dùng không định gọi, tiền không quay lại (hố đen); cờ symexec chỉ là gợi ý",
    "  T3 transfer/transferFrom tới contract token với người nhận = chính contract token đó, tx thành công",
    "  T4 như T1 + cùng địa chỉ có contract ở chain khác tạo TRƯỚC tx (xc_timing = before; ghi xc_chain = origin_net)",
    "  T5 KEY (≥ 2 lần rút nhanh về ví lạ) / ADDR bị rút / 7702 tới delegate sweeper",
    "  T6 address poisoning: đủ 3 dấu hiệu ở CODEBOOK ca 18",
    "  Ưu tiên: T4 > T1; T5 > T6",
    "subtype: NCD / PCD / KEY / ADDR / 7702 / -   (T1/T4: PCD nếu own_create_time có giá trị, tức contract tạo SAU tx; NCD nếu '-'. T2/T3/T6/âm: '-')",
    "xc_chain: chỉ với T4; còn lại '-'",
    "malicious: 1 = đã thấy tiền về ví bên thứ ba; 0 = đã kiểm tra và không có; ? = chưa thấy nhưng có dấu hiệu; - = mẫu âm",
    "confidence: high / medium / low. high chỉ khi đã mở tx_url; chỉ dựa vào cột bằng chứng thì tối đa medium",
    "evidence: bằng chứng đã dùng, ví dụ: code, tx, timing, xc, delegate, history",
    "note: BẮT BUỘC khi confidence = low, is_misuse = ?, hoặc mẫu là ca khó (CODEBOOK §5)",
    "",
    "Quy trình (CODEBOOK §4.1):",
    "0. Mở tx_url (Status, From, To, Value, Method, token chuyển).",
    "A. code_at_event: none → B;  code / destroyed → C;  7702 → D;  unknown → xét theo cột delegate_7702 (hai dòng dưới, chọn đúng MỘT):",
    "   unknown, có delegate_7702 (≠ '-'): delegate hôm nay tạo SAU tx. Xem tab Authorizations: lúc tx không ủy quyền → D2; ủy quyền tới delegate X → D với X (ghi địa chỉ ĐẦY ĐỦ 0x… của X vào note).",
    "   unknown, không có delegate_7702 ('-'): xem tx tạo contract trên explorer rồi chọn B/C/D.",
    "B. calldata_kind ≠ call ('-' chuyển coin thuần, text văn bản, short), hoặc calldata_len = 32 mà Method trên tx là mã thô 0x… → D. calldata_kind = unknown: xem Input trên tx_url.",
    "   calldata_kind = call: xc_timing before → T4;  after → T1;  unknown → mở explorer_url_xc, vẫn không biết → T4 low;  '-' → T1.",
    "C. transfer/transferFrom với người nhận = chính address, thành công → T3.  Contract định gọi, gọi đúng hàm → NONE.",
    "   T2 chỉ khi đủ 3: tx success; contract không phải thứ định gọi; tiền không quay lại. Thiếu một → NONE. Tx fail → NONE.",
    "D. 7702: tra data/gold/delegate_lookup.csv (final_class) và CODEBOOK ca 13; ủy quyền xảy ra SAU tx thì xét như EOA thường.",
    "   forwarder_sweeper → T5/7702, TRỪ KHI ví đích của forwarder là ví có NHÃN sàn trên explorer (ca 7a) → NONE.",
    "   Đếm lần rút nhanh (vào rồi ≤ 10 phút ra ≥ 90% về ví lạ) trong toàn lịch sử: ≥ 2 → T5/KEY;  đúng 1 → NONE, low;  0 → T6 nếu đủ dấu hiệu, còn lại NONE.",
    "   Chỉ rút về ví có NHÃN sàn / custodian trên explorer là KHÔNG tính (ca 7a). Ví nhận không nhãn (ca 7b, 7c) hoặc có nhãn router / bridge (ca 7d) VẪN tính: confidence low, note 'ca 7b' / 'ca 7c' / 'ca 7d'.",
    "",
    "Ca khó hay gặp (CODEBOOK §5):",
    "- selector 0x64617461 là chữ \"data\" (inscription): calldata_kind = text, không phải gọi hàm.",
    "- Cùng giây giữa hai chain = không trước → xc_timing after → T1.",
    "- Cùng contract trên mọi chain kể cả chain này (Multicall3, Permit2, Safe…) → âm. Gọi hàm tới predeploy của chain KHÁC khi chain này không có code → T4.",
    "- 7702 tới ví thông minh hợp lệ, chủ ví vẫn rút bình thường → âm. Xem tiền thực tế đi đâu, không suy ra từ symexec.",
    "- Tx chỉ mất gas hoặc tx lỗi tới địa chỉ không có code mà có gọi hàm vẫn là T1/T4.",
    "Chi tiết, ví dụ đã chốt và quy tắc mặc định khi phân vân: CODEBOOK §4.",
    "",
    "Xong thì lưu file (giữ nguyên tên sheet Gan_nhan), chạy `sample_gold check <file>` và báo người điều phối.",
]


GUIDE_MAIN = [
    "MẪU CHÍNH (CODEBOOK v0.6.2): mỗi dòng là MỘT TX (tx_hash, tx_url) tới address trên chain. Nhãn áp cho TX NÀY.",
    "\"Trước tx\" = trước tx này trên chuỗi (block, rồi vị trí trong block); không dùng ngày hay tx đầu tiên. at_time là giờ UTC của tx.",
    "code_at_event = code ngay trước tx (none / code / destroyed / 7702 / unknown); own_create_time = lúc tạo contract trên chain này, nếu có.",
    "Với EOA (KEY/ADDR/7702) được dùng TOÀN BỘ lịch sử của địa chỉ, kể cả sau tx, để nhận ra sweeper; nhãn vẫn là của tx.",
    "7702: file này chỉ được tạo sau khi A và B đã rà xong bảng delegate (mỗi người một bản riêng: class_X = ok, hoặc lớp khác + lý do ở note_X) và mọi phản đối đã thống nhất. Khi gán dùng final_class.",
    "Thứ tự: làm file 'double' TRƯỚC; file 'single' chỉ làm khi người điều phối báo. Mỗi phiên tối đa ~40 mẫu.",
    "",
]


def cmd_sheet(who: str, which: str = "pilot", only_double: bool = False, only_single: bool = False,
              skip_delegate_check: bool = False) -> None:
    """File Excel cho người gán `who`: bằng chứng (đã làm mù) + cột nhãn có danh sách chọn + sheet hướng dẫn.
    Mẫu chính: từ chối tạo file khi bảng delegate chưa được A và B rà đủ (GOLD_PROTOCOL §2.3). `skip_delegate_check` chỉ cho
    test / chạy thử, in cảnh báo, không dùng để tạo file thật cho người gán."""
    if which == "main":
        errs = check_delegates(pd.read_csv(GOLD_DIR / "delegate_lookup.csv", dtype=str, keep_default_na=False), strict=True)
        if errs and not skip_delegate_check:
            raise SystemExit(f"Không tạo file gán: bảng delegate chưa rà xong hoặc còn phản đối chưa giải quyết ({len(errs)} lỗi). "
                             "A và B phải rà hết, mọi phản đối phải thống nhất (class_A, class_B ∈ {ok, final_class}); "
                             "`sample_gold check-delegates --strict` phải báo 0 lỗi.")
        if errs:
            bar = "!" * 70
            print(f"{bar}\nCẢNH BÁO: BỎ QUA cổng rà delegate ({len(errs)} lỗi). File này CHỈ để thử, "
                  f"KHÔNG giao cho người gán.\n{bar}")
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    items = pd.read_csv(GOLD_DIR / ("pilot50_items.csv" if which == "pilot" else "main_items.csv"), dtype=str).fillna("-")
    if only_double and only_single:
        raise SystemExit("chọn một trong --only-double / --only-single")
    if only_double or only_single:                       # cả hai người làm tập gán đôi trước; tập đơn chỉ làm sau khi kiểm κ 50 mẫu đầu
        key = pd.read_csv(GOLD_DIR / "main_key.csv", dtype=str)
        dbl = key.double == "True"
        items = items[items.item_id.isin(key[dbl if only_double else ~dbl].item_id)].reset_index(drop=True)
    evidence = [c for c in items.columns if c not in ("item_id", "chain", "address")]
    cols = ["item_id", "chain", "address", *LABEL_COLS, *evidence]
    part = "_double" if only_double else "_single" if only_single else ""
    out = GOLD_DIR / (f"annotator_{who}.xlsx" if which == "pilot" else f"annotator_main{part}_{who}.xlsx")
    if out.exists():
        raise SystemExit(f"{out} đã có, không ghi đè (có thể đã có nhãn). Xoá/đổi tên trước nếu muốn tạo lại.")
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(cols)
    for _, r in items.iterrows():
        ws.append([r.get(c, "") for c in cols])
    yellow, bold = PatternFill("solid", fgColor="FFF2CC"), Font(bold=True)
    last = len(items) + 1
    for j, c in enumerate(cols, 1):
        ws.cell(1, j).font = bold
        ws.column_dimensions[get_column_letter(j)].width = 44 if c in ("address", "note") else 12 if c in LABEL_COLS else 16
        for i in range(2, last + 1):
            cell = ws.cell(i, j)
            if c in LABEL_COLS:
                cell.fill = yellow
                cell.number_format = "@"                        # giữ "1"/"0" là chữ, không thành số
            elif (c.startswith("explorer_url") or c == "tx_url") and str(cell.value).startswith("http"):
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"
        if c in LABEL_COLS:
            ws.cell(1, j).fill = yellow
    for c, opts in CHOICES.items():
        col = get_column_letter(cols.index(c) + 1)
        dv = DataValidation(type="list", formula1='"' + ",".join(opts) + '"', allow_blank=True)
        dv.add(f"{col}2:{col}{last}")
        ws.add_data_validation(dv)
    ws.freeze_panes = "D2"                                      # giữ dòng tiêu đề + mã, chain, địa chỉ
    g = wb.create_sheet("Huong_dan")
    for line in (GUIDE_MAIN if which == "main" else []) + GUIDE:
        g.append([line])
    g["A1"].font = bold
    g.column_dimensions["A"].width = 140
    for row in g.iter_rows():
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(out)
    print(f"Ghi {len(items)} mẫu → {out}")


def read_labels(path: str) -> pd.DataFrame:
    """Nhãn của một người gán, từ .csv (annotator_template) hoặc .xlsx (`cmd_sheet`)."""
    if str(path).lower().endswith(".xlsx"):
        d = pd.read_excel(path, sheet_name=SHEET, dtype=str)
    else:
        d = pd.read_csv(path, dtype=str)
    d = d[["item_id", *[c for c in ("address", *LABEL_COLS) if c in d.columns]]].fillna("")
    return d.apply(lambda s: s.str.strip())


def check_labels(d: pd.DataFrame) -> list[str]:
    """Lỗi hình thức theo CODEBOOK §2 (không so với nhãn tác giả hay người gán kia)."""
    errs = []
    for r in d.itertuples(index=False):
        g = r._asdict()
        at = lambda c: g.get(c, "")
        for c, opts in CHOICES.items():
            v = at(c)
            if c in ("is_misuse", "type", "confidence") and not v:
                errs.append(f"{r.item_id}: thiếu {c}")
            elif v and v.lower() not in [o.lower() for o in opts]:
                errs.append(f"{r.item_id}: {c} = {v!r} không hợp lệ")
        m, t = at("is_misuse"), at("type").upper()
        if m == "0" and t and t != "NONE":
            errs.append(f"{r.item_id}: is_misuse = 0 nhưng type = {t} (phải NONE)")
        if m == "1" and t == "NONE":
            errs.append(f"{r.item_id}: is_misuse = 1 nhưng type = NONE")
        xc = at("xc_chain")
        if t == "T4" and xc in ("", "-"):
            errs.append(f"{r.item_id}: T4 nhưng chưa ghi xc_chain")
        if t != "T4" and xc not in ("", "-"):
            errs.append(f"{r.item_id}: xc_chain = {xc} chỉ dùng với T4 (ghi vào note nếu cần)")
        if (at("confidence").lower() == "low" or m == "?") and not at("note"):
            errs.append(f"{r.item_id}: confidence low hoặc is_misuse ? thì phải có note")
    n_q = int((d.is_misuse == "?").sum())
    if n_q > len(d) // 10:
        errs.append(f"có {n_q} mẫu is_misuse = ? (tối đa {len(d) // 10})")
    return errs


def cmd_check(path: str) -> None:
    d = read_labels(path)
    errs = check_labels(d)
    filled = d.is_misuse != ""
    print(f"{path}: {int(filled.sum())}/{len(d)} mẫu đã điền is_misuse")
    for col in ("is_misuse", "type", "confidence"):
        print(f"  {col}: " + ", ".join(f"{k or '(trống)'} {n}" for k, n in d[col].value_counts().items()))
    print("Không có lỗi hình thức." if not errs else f"{len(errs)} lỗi:\n  " + "\n  ".join(errs))


# ---------------------------------------------------------------- κ

def cohen_kappa(a: list, b: list) -> float:
    """Cohen's κ cho hai dãy nhãn danh định cùng độ dài."""
    if len(a) != len(b) or not a:
        raise ValueError("hai dãy nhãn phải cùng độ dài, khác rỗng")
    cats = sorted(set(a) | set(b), key=str)
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    pe = sum((a.count(c) / n) * (b.count(c) / n) for c in cats)
    return 1.0 if pe == 1 else (po - pe) / (1 - pe)


def pabak(a: list, b: list) -> float:
    """Prevalence-adjusted bias-adjusted κ = 2·po − 1 (hai lớp). κ thấp giả khi tỉ lệ dương lệch xa 50%."""
    return 2 * sum(x == y for x, y in zip(a, b)) / len(a) - 1


def gwet_ac1(a: list, b: list) -> float:
    """Gwet's AC1 cho hai người gán, nhãn danh định: bền hơn κ khi mẫu lệch lớp."""
    cats = sorted(set(a) | set(b), key=str)
    n, q = len(a), len(cats)
    po = sum(x == y for x, y in zip(a, b)) / n
    pi = [(a.count(c) + b.count(c)) / (2 * n) for c in cats]
    pe = sum(p * (1 - p) for p in pi) / (q - 1) if q > 1 else 0.0
    return 1.0 if pe == 1 else (po - pe) / (1 - pe)


def cluster_boot_ci(a: list, b: list, clusters: list, stat=cohen_kappa, boot: int = 2000, seed: int = 0) -> list[float]:
    """CI 95% (percentile) của `stat` khi lấy lại mẫu THEO CỤM (ví dụ cùng delegate 7702 hoặc cùng deployer): mẫu cùng cụm không độc lập."""
    rng = np.random.default_rng(seed)
    codes = pd.factorize(pd.Series(clusters))[0]
    idx = [np.flatnonzero(codes == c) for c in range(codes.max() + 1)]
    a, b = np.array(a, dtype=object), np.array(b, dtype=object)
    out = []
    for _ in range(boot):
        i = np.concatenate([idx[j] for j in rng.integers(0, len(idx), len(idx))])
        out.append(stat(a[i].tolist(), b[i].tolist()))
    return [float(np.quantile(out, .025)), float(np.quantile(out, .975))]


def cmd_kappa(path_a: str, path_b: str, items_csv: str | None = None, key_csv: str | None = None) -> dict[str, float]:
    a = read_labels(path_a)
    b = read_labels(path_b)
    j = a.merge(b, on="item_id", suffixes=("_A", "_B"))
    if "address_A" in j and "address_B" in j:                   # hai người phải gán cùng một bộ mẫu
        diff = j[j.address_A.str.lower() != j.address_B.str.lower()]
        if len(diff):
            raise SystemExit(f"{len(diff)} item_id trỏ tới địa chỉ khác nhau, ví dụ {diff.item_id.iloc[0]}: hai file không cùng bộ mẫu")
    j = j[(j.is_misuse_A != "") & (j.is_misuse_B != "")].reset_index(drop=True)
    clusters = list(range(len(j)))                               # mặc định: từng mẫu là một cụm
    if items_csv:                     # cụm = delegate 7702, rồi người tạo contract trên chain này, rồi deployer ở mạng khác, rồi từng mẫu
        it = pd.read_csv(items_csv, dtype=str).set_index("item_id")
        def cl(i):
            r = it.loc[i] if i in it.index else None
            for c in ("delegate_7702", "creator", "deployer"):
                if r is not None and str(r.get(c, "-")) not in ("-", "nan"):
                    return f"{c}:{r[c]}"
            return f"item:{i}"
        clusters = [cl(i) for i in j.item_id]
    out = {}
    for col in ("is_misuse", "type"):
        x = j[f"{col}_A"].str.strip().str.upper().tolist()
        y = j[f"{col}_B"].str.strip().str.upper().tolist()
        out[col] = cohen_kappa(x, y)
        ci = cluster_boot_ci(x, y, clusters)
        print(f"κ({col}) = {out[col]:.3f} [CI95 theo cụm {ci[0]:.3f}–{ci[1]:.3f}] trên {len(j)} mẫu, đồng ý {sum(p == q for p, q in zip(x, y))}/{len(j)}"
              f"; Gwet AC1 {gwet_ac1(x, y):.3f}" + (f"; PABAK {pabak(x, y):.3f}" if col == "is_misuse" else ""))
    from address_misuse.gold_eval import ca7_unlabelled, delegate_decided, delegate_named
    subsets = []                      # tiền đăng ký: κ khi bỏ một nhóm mẫu (GOLD_PROTOCOL §4 j, j′, k)
    if items_csv:
        dd = set(delegate_decided(it.reset_index()))
        subsets.append(("no_delegate", "do delegate quyết định", dd))
        lk = GOLD_DIR / "delegate_lookup.csv"
        lookup = pd.read_csv(lk, dtype=str, keep_default_na=False) if lk.exists() else pd.DataFrame(columns=["chain", "delegate"])
        named = set()
        for w in ("A", "B"):          # (j′) mẫu unknown mà một trong hai người ghi delegate (có trong bảng) vào note
            if f"note_{w}" in j:
                named |= set(delegate_named(j.rename(columns={f"note_{w}": "note"}), it.reset_index(), lookup))
        subsets.append(("no_delegate_any", "do delegate quyết định, kể cả unknown có delegate trong bảng ở note (j′)", dd | named))
    ca7 = set(ca7_unlabelled(j.rename(columns={"note_A": "note"}))) | set(ca7_unlabelled(j.rename(columns={"note_B": "note"})))
    subsets.append(("no_ca7", "có note ca 7b/7c/7d ở một trong hai người", ca7))
    for name, what, drop in subsets:
        keep = ~j.item_id.isin(drop).to_numpy()
        for col in ("is_misuse", "type"):
            x = [v for v, k in zip(j[f"{col}_A"].str.strip().str.upper(), keep) if k]
            y = [v for v, k in zip(j[f"{col}_B"].str.strip().str.upper(), keep) if k]
            if x:
                out[f"{col}_{name}"] = cohen_kappa(x, y)
                ci = cluster_boot_ci(x, y, [c for c, k in zip(clusters, keep) if k])
                print(f"  bỏ {int((~keep).sum())} mẫu {what}: κ({col}) = {out[f'{col}_{name}']:.3f} "
                      f"[CI95 theo cụm {ci[0]:.3f}–{ci[1]:.3f}] trên {len(x)} mẫu")
    q = (j.is_misuse_A.str.upper() == "?") | (j.is_misuse_B.str.upper() == "?")
    if q.any():
        x, y = j.loc[~q, "is_misuse_A"].str.upper().tolist(), j.loc[~q, "is_misuse_B"].str.upper().tolist()
        print(f"  bỏ {int(q.sum())} mẫu có '?': κ(is_misuse) = {cohen_kappa(x, y):.3f} trên {len(x)} mẫu")
    if key_csv:                                                  # % đồng ý theo tầng (chỉ người thiết kế mới được mở file khoá)
        k = pd.read_csv(key_csv, dtype=str)[["item_id", "stratum"]]
        g = j.merge(k, on="item_id")
        g["same"] = g.is_misuse_A.str.upper() == g.is_misuse_B.str.upper()
        print("  % đồng ý is_misuse theo tầng: " + ", ".join(f"{s} {m.mean():.0%} ({len(m)})" for s, m in g.groupby("stratum").same))
    dis = j[(j.is_misuse_A.str.upper() != j.is_misuse_B.str.upper()) | (j.type_A.str.upper() != j.type_B.str.upper())]
    if len(dis):
        print("\nBất đồng cần phân xử:")
        print(dis[["item_id", "is_misuse_A", "type_A", "is_misuse_B", "type_B"]].to_string(index=False))
    return out


def main():
    ap = argparse.ArgumentParser()
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("pilot")
    p.add_argument("--seed", type=int, default=0)
    c = sub.add_parser("check")
    c.add_argument("path")
    sh = sub.add_parser("sheet")
    sh.add_argument("who", nargs="?", default="B")
    sh.add_argument("--set", choices=["pilot", "main"], default="pilot")
    sh.add_argument("--only-double", action="store_true")
    sh.add_argument("--only-single", action="store_true")
    sh.add_argument("--skip-delegate-check", action="store_true",
                    help="CHỈ cho test / chạy thử: bỏ cổng rà bảng delegate. KHÔNG dùng để tạo file thật cho người gán")
    mn = sub.add_parser("main")
    mn.add_argument("--seed", type=int, default=0)
    mn.add_argument("--double", type=int, default=200)
    mn.add_argument("--live-from", help="main_items.csv cũ: dùng lại ảnh chụp getCode/symexec hôm nay thay vì tra lại")
    cd = sub.add_parser("check-delegates")
    cd.add_argument("path", nargs="?")
    cd.add_argument("--who", choices=["A", "B"], help="kiểm bản rà riêng của một người")
    cd.add_argument("--strict", action="store_true",
                    help="cổng trước khi tạo file gán: mọi phản đối phải đã giải quyết (class_X ∈ {ok, final_class})")
    dc = sub.add_parser("delegate-copy", help="tạo bản rà riêng delegate_review_<A|B>.csv")
    dc.add_argument("who", choices=["A", "B"])
    dc.add_argument("--out")
    dm = sub.add_parser("merge-delegates", help="gộp hai bản rà riêng vào delegate_lookup.csv")
    dm.add_argument("a")
    dm.add_argument("b")
    k = sub.add_parser("kappa")
    k.add_argument("a")
    k.add_argument("b")
    k.add_argument("--items", help="main_items.csv: bootstrap CI theo cụm delegate/deployer")
    k.add_argument("--key", help="main_key.csv: thêm % đồng ý theo tầng (chỉ người thiết kế)")
    args = ap.parse_args()
    if args.cmd == "pilot":
        cmd_pilot(args.seed)
    elif args.cmd == "check":
        cmd_check(args.path)
    elif args.cmd == "main":
        cmd_main(args.seed, args.double, args.live_from)
    elif args.cmd == "check-delegates":
        cmd_check_delegates(args.path, args.who, strict=args.strict)
    elif args.cmd == "delegate-copy":
        cmd_delegate_copy(args.who, args.out)
    elif args.cmd == "merge-delegates":
        cmd_merge_delegates(args.a, args.b)
    elif args.cmd == "sheet":
        cmd_sheet(args.who, args.set, args.only_double, args.only_single, args.skip_delegate_check)
    else:
        cmd_kappa(args.a, args.b, args.items, args.key)


if __name__ == "__main__":
    main()
